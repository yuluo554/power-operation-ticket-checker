"""docx 渲染（extras[report]：python-docx）：章节结构 → .docx 文件/字节流。

- 懒加载 python-docx，缺失时抛 ReportUnavailable（CLI 打印安装提示）；
- 无外链纪律：不写超链接、不写 URL（判据 find_external_links，测试锁定）；
- 渲染仅搬运 build_report_sections 的章节结构，不产生新事实。
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, Optional

from . import ReportUnavailable
from .sections import build_report_sections


def ensure_docx_available() -> None:
    """前置检查：python-docx 不可用即抛 ReportUnavailable（批量任务开始前先探测，避免半途而废）。"""
    try:
        import docx  # noqa: F401
    except ImportError as e:
        raise ReportUnavailable(f"python-docx 未安装（{e}）")


def _load_document():
    try:
        from docx import Document
        from docx.oxml.ns import qn

        return Document, qn
    except ImportError as e:
        raise ReportUnavailable(f"python-docx 未安装（{e}）")


def _apply_cjk_font(doc, qn, font_name: str = "宋体") -> None:
    """Normal 样式设置中文字体（失败仅影响观感，不影响内容，不抛）。"""
    try:
        style = doc.styles["Normal"]
        style.font.name = font_name
        rpr = style.element.rPr
        if rpr is not None and rpr.rFonts is not None:
            rpr.rFonts.set(qn("w:eastAsia"), font_name)
    except Exception:
        pass


def _render_section(doc, section: Dict[str, Any]) -> None:
    kind = section["kind"]
    if kind == "title":
        doc.add_heading(section["text"], level=0)
    elif kind == "meta_table":
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in section["rows"]:
            cells = table.add_row().cells
            cells[0].text = str(label)
            cells[1].text = str(value)
            for p in cells[0].paragraphs:
                for r in p.runs:
                    r.font.bold = True
    elif kind == "heading":
        doc.add_heading(section["text"], level=section.get("level", 1))
    elif kind == "paragraph":
        doc.add_paragraph(section["text"])
    elif kind == "bullets":
        for item in section["items"]:
            doc.add_paragraph(str(item), style="List Bullet")
    elif kind == "table":
        header = [str(h) for h in section["header"]]
        table = doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        for i, h in enumerate(header):
            table.rows[0].cells[i].text = h
            for p in table.rows[0].cells[i].paragraphs:
                for r in p.runs:
                    r.font.bold = True
        for row_values in section["rows"]:
            cells = table.add_row().cells
            for i, value in enumerate(row_values):
                cells[i].text = "" if value is None else str(value)
    else:  # 未知章节类型：正文文本兜底，不静默丢弃
        doc.add_paragraph(str(section))


def render_document(result: dict, meta: Optional[dict] = None):
    """run_pipeline 结果 → python-docx Document（供高级用例继续加工）。"""
    Document, qn = _load_document()
    doc = Document()
    _apply_cjk_font(doc, qn)
    for section in build_report_sections(result, meta):
        _render_section(doc, section)
    return doc


def render_docx_report(result: dict, out_path, meta: Optional[dict] = None) -> Path:
    """渲染并保存 docx 到 out_path（父目录自动创建），返回实际路径。"""
    doc = render_document(result, meta)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def render_docx_report_bytes(result: dict, meta: Optional[dict] = None) -> bytes:
    """渲染为 docx 字节流（Web 下载用）。"""
    doc = render_document(result, meta)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
